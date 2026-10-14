#!/usr/bin/env python3
"""Verify evidence links, referential integrity, author order, and workbook contents."""
import json
import math
import argparse
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from export import ROOT, NS, build_tables, metadata_tables


def verify_reader(openpyxl, path, tables):
    workbook = openpyxl.load_workbook(path, read_only=False, data_only=False)
    assert workbook.sheetnames == list(tables)
    for name, expected in tables.items():
        worksheet = workbook[name]
        assert worksheet.max_row == len(expected) + 1
        assert worksheet.freeze_panes == "A2"
        assert worksheet.auto_filter.ref
        if expected:
            headers = [cell.value for cell in worksheet[1]]
            observed = list(worksheet.iter_rows(min_row=2, values_only=True))
            for row, original in zip(observed, expected):
                for header, value in zip(headers, row):
                    wanted = original.get(header, "")
                    if name == "Readme" and original.get("topic") == "Generated" and header == "description":
                        datetime.fromisoformat(value)
                    elif wanted in ("", None):
                        assert value in ("", None)
                    elif isinstance(wanted, float):
                        assert math.isclose(value, wanted)
                    else:
                        assert value == wanted, f"Workbook value mismatch in {name}.{header}"
    workbook.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--excel-reader-path", type=Path, help="Optional directory containing openpyxl for an independent workbook-reader check")
    args = parser.parse_args()
    tables = build_tables()
    papers = tables["Publications"]
    ids = {p["paper_id"] for p in papers}
    assert len(ids) == len(papers), "Duplicate paper IDs"
    assert len({(p["venue"], p["year"], p["title"].casefold()) for p in papers}) == len(papers)
    assert papers, "No publications collected"
    if "Proceedings_Ledger" in tables:
        ledger = tables["Proceedings_Ledger"]
        assert len({p["ledger_id"] for p in ledger}) == len(ledger)
        linked_ids = [p["project_paper_id"] for p in ledger if p["project_paper_id"]]
        assert len(linked_ids) == len(set(linked_ids)) and set(linked_ids) == ids
        cells = tables["Completeness"]
        assert {(c["venue"], c["year"]) for c in cells} == {(v, y) for v in ["NSDI", "SIGCOMM", "CoNEXT", "IMC", "HotNets"] for y in range(2019, 2027)}
        assert sum(c["inventory_records"] for c in cells) == len(ledger)
        assert tables["Unresolved_Queue"] == [p for p in ledger if p["decision"] == "unresolved"]
        for entry in ledger:
            for cache in entry["evidence_caches"].split("; "):
                assert not cache or (ROOT / cache).exists(), f"Missing ledger evidence: {cache}"
    author_rows = defaultdict(list)
    for row in tables["Author_Affiliations"]:
        assert row["paper_id"] in ids
        assert row["source_url"] and (ROOT / row["source_cache"]).exists()
        author_rows[row["paper_id"]].append(row)
    for p in papers:
        assert p["title"] and p["url"] and p["year"]
        assert (ROOT / p["source_cache"]).exists()
        assert p["abstract_source"] or not p["abstract"]
        rows = author_rows[p["paper_id"]]
        positions = {r["author_position"] for r in rows}
        assert positions == set(range(1, p["author_count"] + 1))
        assert p["author_count"] > 0
        assert any(r["is_tencent"] for r in rows), f'Tencent mapping missing: {p["title"]}'
        for position in positions:
            assert len({r["author_name"] for r in rows if r["author_position"] == position}) == 1
    weights = defaultdict(float)
    seen_edges = set()
    institutions = {i["institution_id"] for i in tables["Institutions"]}
    assert "" not in institutions, "Empty institution entity"
    for edge in tables["Collaboration"]:
        assert edge["paper_id"] in ids and edge["partner_institution_id"] in institutions
        key = (edge["paper_id"], edge["partner_institution_id"])
        assert key not in seen_edges
        seen_edges.add(key)
        weights[edge["paper_id"]] += edge["fractional_partner_weight"]
    assert all(math.isclose(w, 1) for w in weights.values())
    for i in tables["Institutions"]:
        assert not i["country"] or i["geography_source"], "Geography lacks source"
    path = ROOT / "output/tencent_networking_publications.xlsx"
    with ZipFile(path) as archive:
        assert archive.testzip() is None
        for name in archive.namelist():
            ET.fromstring(archive.read(name))
        for index, (name, expected) in enumerate(tables.items(), 1):
            sheet = ET.fromstring(archive.read(f"xl/worksheets/sheet{index}.xml"))
            rows = sheet.findall(f"{{{NS}}}sheetData/{{{NS}}}row")
            assert len(rows) == len(expected) + 1, f"Wrong row count: {name}"
            assert not sheet.findall(f".//{{{NS}}}f"), "Unexpected formula"
    reader_checked = False
    if args.excel_reader_path:
        sys.path.insert(0, str(args.excel_reader_path.resolve()))
        import openpyxl
        verify_reader(openpyxl, path, tables)
        if "Metadata_Checks" in tables:
            compact = metadata_tables(tables)
            assert len(compact["Publications"]) == 32
            assert all(p["abstract_status"].startswith("complete_") and p["metadata_status"] == "required_metadata_complete" for p in compact["Publications"])
            verify_reader(openpyxl, ROOT / "output/tencent_networking_metadata.xlsx", compact)
        reader_checked = True
    result = {"status": "passed", "publications": len(papers), "author_affiliation_rows": len(tables["Author_Affiliations"]),
              "collaboration_edges": len(seen_edges), "independent_excel_reader_checked": reader_checked,
              "note": "Structural and independent-reader validation. User-reviewed primary metadata is complete; full venue recall, version families, person identities and remaining publisher PDF access are not certified."}
    (ROOT / "output/validation_report.json").write_text(json.dumps(result, indent=2) + "\n")
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
